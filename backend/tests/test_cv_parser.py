from datetime import datetime

from io import BytesIO

from docx import Document
from reportlab.pdfgen import canvas

from app.services.cv_parsing import clean_reviewed_profile, parse_cv_document
from app.services.rule_based import parse_candidate_from_cv
from app.services.parser import CVTextParser
from app.services import cv_fields


SAMPLE_CV = """
NGUYEN VAN AN
Senior Backend Engineer
Ho Chi Minh City, Vietnam
an.nguyen@example.com | +84 912 345 678
linkedin.com/in/nguyenvanan | https://github.com/nguyenvanan

PROFESSIONAL SUMMARY
Backend engineer building payment platforms and distributed systems for more than 7 years.

TECHNICAL SKILLS
Python, FastAPI, PostgreSQL, Redis, Docker, Kubernetes, AWS, Kafka, Team leadership

WORK EXPERIENCE
Senior Backend Engineer | FinPay JSC
2021 - Present
- Improved payment throughput by 45% and led a team of 6 engineers.
Backend Engineer | ShopNow Ltd
2018 - 2021
- Built ecommerce APIs with Python and PostgreSQL.

PROJECTS
Fraud detection platform | Python, machine learning, Kafka
Merchant settlement service | FastAPI, PostgreSQL

EDUCATION
Bachelor of Computer Science | University of Technology | 2014 - 2018

CERTIFICATIONS
AWS Certified Solutions Architect

LANGUAGES
Vietnamese - Native | English - Professional
"""


def test_parser_extracts_full_candidate_profile():
    parsed = parse_candidate_from_cv(SAMPLE_CV)

    assert parsed["name"] == "NGUYEN VAN AN"
    assert parsed["email"] == "an.nguyen@example.com"
    assert parsed["current_title"] == "Senior Backend Engineer"
    assert parsed["years_of_experience"] == 7
    assert parsed["previous_companies"] == ["FinPay JSC", "ShopNow Ltd"]
    assert parsed["experience_timeline"][0]["company"] == "FinPay JSC"
    assert "python" in parsed["skills"]
    assert "github" not in parsed["skills"]
    assert len(parsed["projects"]) == 2
    assert parsed["achievements"]
    assert parsed["completeness_score"] >= 90
    assert parsed["missing_critical_fields"] == []


def test_reviewed_profile_preserves_preferences_without_provider_metadata():
    parsed = clean_reviewed_profile({
        "name": "Alice Nguyen", "preferred_location": "  Da Nang  ",
        "notice_period": "30 days", "skills": ["Python", "Python"],
        "rich_text": {"summary": '<p onclick="bad()"><strong>Good</strong><script>alert(1)</script></p>', "owner": "bad"},
        "ai_provider": "old-provider", "gender": "unsupported field",
    })
    assert parsed["preferred_location"] == "Da Nang"
    assert parsed["notice_period"] == "30 days"
    assert parsed["skills"] == ["Python"]
    assert "ai_provider" not in parsed
    assert "gender" not in parsed
    assert parsed["rich_text"] == {"summary": "<p><strong>Good</strong></p>"}


def test_overlapping_roles_are_not_double_counted():
    parsed = parse_candidate_from_cv(
        """
        LE THI BINH
        Software Engineer
        WORK EXPERIENCE
        Lead Engineer | Company A
        2020 - Present
        Software Engineer | Company B
        2021 - 2023
        """
    )

    # The second role sits inside the first interval.
    assert parsed["years_of_experience"] == datetime.now().year - 2020


def test_vietnamese_and_multiline_preferences_are_extracted():
    parsed = parse_candidate_from_cv("""
        NGUYỄN THỊ LINH
        Địa điểm làm việc mong muốn:
        Đà Nẵng, Hà Nội
        Thời gian báo trước: 30 ngày
        """)
    assert parsed["preferred_location"] == "Đà Nẵng, Hà Nội"
    assert parsed["notice_period"] == "30 ngày"
    assert parsed["confidence"]["notice_period"] == "high"


def test_inline_preferences_and_notice_do_not_swallow_each_other():
    parsed = parse_candidate_from_cv("Preferred location: Singapore | Notice period: 2 weeks")
    assert parsed["preferred_location"] == "Singapore"
    assert parsed["notice_period"] == "2 weeks"


def test_unknown_preferences_stay_empty_without_using_home_address():
    parsed = parse_candidate_from_cv("LE THI BINH\nLocation: Hanoi\nSoftware Engineer\nEducation\nUniversity 2015 - 2019")
    assert parsed["preferred_location"] is None
    assert parsed["notice_period"] is None
    assert parsed["years_of_experience"] is None


def test_numeric_month_dates_count_only_professional_experience():
    parsed = parse_candidate_from_cv("""
        LE THI BINH
        Software Engineer
        Age: 30 years
        Work experience
        Engineer | Company A
        01/2020 - 12/2022
        Engineer | Company B
        2021/06 - 2023/12
        Education
        University 2010 - 2019
        """)
    assert parsed["experience_months"] == 48
    assert parsed["years_of_experience"] == 4
    assert parsed["previous_companies"] == ["Company A", "Company B"]


def test_docx_tables_keep_order_and_local_preferences():
    document = Document()
    document.add_paragraph("TRAN VAN NAM")
    document.add_paragraph("Work experience")
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Engineer | Acme Ltd\nJan 2020 - Dec 2022"
    table.cell(1, 0).text = "Improved response time by 40%."
    document.add_paragraph("Education")
    document.add_paragraph("Bachelor of Computer Science")
    document.add_paragraph("Preferred location: Da Nang")
    document.add_paragraph("Available immediately")
    buffer = BytesIO()
    document.save(buffer)
    parsed, text = parse_cv_document("cv.docx", buffer.getvalue())
    assert text.index("Acme Ltd") < text.index("Education")
    assert parsed["previous_companies"] == ["Acme Ltd"]
    assert parsed["preferred_location"] == "Da Nang"
    assert parsed["notice_period"] == "Immediately"
    assert parsed["source"] == "local_cv_parser"
    assert "ai_provider" not in parsed


def test_notice_duration_in_sentence_is_explicitly_extracted():
    parsed = parse_candidate_from_cv("LE THI BINH\nI am required to give 30 days notice before joining.")
    assert parsed["notice_period"] == "30 days"


def test_future_roles_do_not_create_negative_experience(monkeypatch):
    class FixedTime:
        @staticmethod
        def now():
            return datetime(2026, 1, 10)
    monkeypatch.setattr(cv_fields, "datetime", FixedTime)
    assert cv_fields.employment_months("Aug 2026 - Dec 2026") is None


def test_mixed_pdf_extracts_scanned_page_and_hidden_profile_link(monkeypatch):
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer)
    for index, line in enumerate([
        "LE THI BINH", "Software Engineer", "binh@example.com", "Skills: Python FastAPI",
        "Summary: " + "Software platform experience " * 15,
    ]):
        pdf.drawString(30, 800 - 20 * index, line)
    pdf.linkURL("https://linkedin.com/in/le-thi-binh", (30, 600, 130, 615))
    pdf.showPage()
    pdf.showPage()
    pdf.save()
    ocr_pages = []
    def local_ocr(_content, page_number):
        ocr_pages.append(page_number)
        return "Preferred location: Singapore\nNotice period: 2 weeks"
    monkeypatch.setattr(CVTextParser, "_parse_pdf_with_ocr", local_ocr)
    parsed, _ = parse_cv_document("cv.pdf", buffer.getvalue())
    assert ocr_pages == [2]
    assert parsed["preferred_location"] == "Singapore"
    assert parsed["notice_period"] == "2 weeks"
    assert parsed["linkedin_url"] == "https://linkedin.com/in/le-thi-binh"
