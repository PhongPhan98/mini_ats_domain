from datetime import datetime
from io import BytesIO
import json

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base, get_db
from app.models import Candidate, Job, Organization, User
from app.rbac import get_current_user
from app.routers import applications, candidates, jobs
from app.services.applications import create_application


@pytest.fixture
def workspace(monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine)()
    org = Organization(name="Test", slug="test")
    other_org = Organization(name="Other", slug="other")
    db.add_all([org, other_org])
    db.flush()
    actor = User(organization_id=org.id, email="hr@example.com", full_name="HR", role="admin")
    db.add(actor)
    db.flush()
    app = FastAPI()
    app.include_router(candidates.router)
    app.include_router(applications.router)
    app.include_router(jobs.router)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: actor
    monkeypatch.setattr(candidates.storage, "save_bytes", lambda *_: "local://test.docx")
    monkeypatch.setattr(applications, "run_stage_change_automations", lambda **_: None)
    monkeypatch.setattr(applications, "log_event", lambda *_: None)
    with TestClient(app) as client:
        yield client, db, actor, org, other_org
    db.close()
    engine.dispose()


def test_all_jobs_cards_move_each_application_independently(workspace):
    client, db, actor, org, other_org = workspace
    candidate = Candidate(organization_id=org.id, owner_user_id=actor.id, name="Linh", email="linh@example.com", skills=["python", "fastapi"])
    unassigned = Candidate(organization_id=org.id, owner_user_id=actor.id, name="Nam")
    foreign = Candidate(organization_id=other_org.id, name="Private")
    jobs = [Job(organization_id=org.id, owner_user_id=actor.id, title=title, requirements="Python") for title in ["Backend", "Platform"]]
    db.add_all([candidate, unassigned, foreign, *jobs])
    db.flush()
    first = create_application(db, candidate=candidate, job=jobs[0], owner_user_id=actor.id)
    second = create_application(db, candidate=candidate, job=jobs[1], owner_user_id=actor.id)
    db.commit()

    response = client.get("/api/candidates/pipeline")
    assert response.status_code == 200
    cards = response.json()
    assert len(cards) == 3
    assert len({card["board_key"] for card in cards}) == 3
    first_card = next(card for card in cards if card.get("application_id") == first.id)
    assert first_card["skills"] == ["python", "fastapi"]
    assert first_card["job_title"] == "Backend"
    moved = client.patch(f"/api/applications/{first.id}/stage", json={"stage": "interview"})
    assert moved.status_code == 200
    refreshed = client.get("/api/candidates/pipeline").json()
    assert next(card for card in refreshed if card.get("application_id") == first.id)["status"] == "interview"
    assert next(card for card in refreshed if card.get("application_id") == second.id)["status"] == "applied"
    job_cards = client.get(f"/api/candidates/pipeline?job_id={jobs[0].id}").json()
    assert len(job_cards) == 1
    assert job_cards[0]["skills"] == ["python", "fastapi"]
    assert client.patch(f"/api/candidates/{unassigned.id}/stage", json={"stage": "screening"}).status_code == 200
    assert client.patch(f"/api/candidates/{candidate.id}/stage", json={"stage": "hired"}).status_code == 409
    assert client.patch(f"/api/candidates/{unassigned.id}/stage", json={"stage": "hired", "application_id": 9999}).status_code == 404


def test_preview_review_import_preserves_availability_without_reparsing(workspace, monkeypatch):
    client, db, actor, org, _ = workspace
    document = Document()
    for line in ["NGUYỄN THỊ LINH", "linh@example.com", "Kỹ năng: Python, FastAPI", "Địa điểm làm việc mong muốn: Đà Nẵng", "Thời gian báo trước: 30 ngày"]:
        document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    cv = buffer.getvalue()
    file = ("cv.docx", cv, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    preview = client.post("/api/candidates/parse", files={"file": file})
    assert preview.status_code == 200
    parsed = preview.json()["parsed"]
    assert parsed["preferred_location"] == "Đà Nẵng"
    assert parsed["notice_period"] == "30 ngày"
    assert "ai_provider" not in parsed
    parsed["notice_period"] = "14 days"
    parsed["rich_text"] = {
        "summary": '<h2 style="text-align: center" onclick="bad()">Profile</h2><p><strong>Backend engineer</strong><script>alert(1)</script></p>',
        "forbidden": "discard me",
    }
    monkeypatch.setattr(candidates, "parse_cv_document", lambda *_: pytest.fail("Reviewed imports must not parse again"))
    imported = client.post("/api/candidates/upload", files={"file": file}, data={"reviewed": "true", "edited_json": json.dumps(parsed)})
    assert imported.status_code == 200
    profile = client.get(f'/api/candidates/{imported.json()["id"]}').json()
    assert profile["parsed_json"]["preferred_location"] == "Đà Nẵng"
    assert profile["parsed_json"]["notice_period"] == "14 days"
    assert profile["parsed_json"]["rich_text"] == {
        "summary": '<h2 style="text-align: center">Profile</h2><p><strong>Backend engineer</strong></p>'
    }
    assert profile["skills"] == ["fastapi", "python"]

    updated = client.patch(f'/api/candidates/{profile["id"]}', json={
        "summary": "Updated profile",
        "rich_text": {"summary": '<p style="background-color: #fff2a8" onmouseover="bad()"><u>Updated profile</u></p><iframe>bad</iframe>'},
    })
    assert updated.status_code == 200
    assert updated.json()["summary"] == "Updated profile"
    assert updated.json()["parsed_json"]["rich_text"]["summary"] == '<p style="background-color: #fff2a8"><u>Updated profile</u></p>'


def test_pipeline_interviewer_has_no_move_controls(workspace):
    client, db, actor, org, _ = workspace
    actor.role = "interviewer"
    db.add(Candidate(organization_id=org.id, owner_user_id=actor.id, name="Alice"))
    db.commit()
    cards = client.get("/api/candidates/pipeline").json()
    assert cards[0]["can_move"] is False
    assert client.patch(f'/api/candidates/{cards[0]["id"]}/stage', json={"stage": "interview"}).status_code == 403


def test_reimport_preserves_preferences_and_existing_ownership(workspace):
    client, db, actor, org, _ = workspace
    candidate = Candidate(organization_id=org.id, owner_user_id=actor.id,
        name="Linh", email="linh@example.com", parsed_json={
            "preferred_location": "Da Nang", "notice_period": "30 days",
            "owner_email": "original@example.com", "owner_user_id": actor.id,
        })
    db.add(candidate)
    db.commit()
    document = Document()
    document.add_paragraph("Linh Nguyen\nlinh@example.com\nSkills: Python, FastAPI")
    buffer = BytesIO()
    document.save(buffer)
    response = client.post("/api/candidates/upload", files={"file": ("cv.docx", buffer.getvalue())})
    assert response.status_code == 200
    profile = response.json()
    assert profile["id"] == candidate.id
    assert profile["parsed_json"]["preferred_location"] == "Da Nang"
    assert profile["parsed_json"]["notice_period"] == "30 days"
    assert profile["parsed_json"]["owner_email"] == "original@example.com"
    assert "python" in profile["skills"]


def test_deleted_job_application_is_not_a_movable_unassigned_candidate(workspace):
    client, db, actor, org, _ = workspace
    candidate = Candidate(organization_id=org.id, owner_user_id=actor.id, name="Linh")
    job = Job(organization_id=org.id, owner_user_id=actor.id, title="Backend", requirements="Python", deleted_at=datetime.utcnow())
    db.add_all([candidate, job])
    db.flush()
    create_application(db, candidate=candidate, job=job, owner_user_id=actor.id)
    db.commit()
    assert client.get("/api/candidates/pipeline").json() == []


def test_job_matching_runs_locally_and_records_local_metadata(workspace):
    client, db, actor, org, _ = workspace
    candidate = Candidate(organization_id=org.id, owner_user_id=actor.id,
        name="Linh", skills=["python", "fastapi"], years_of_experience=5,
        summary="Backend engineer Python FastAPI 3 years of experience",
        parsed_json={"current_title": "Backend Engineer"})
    job = Job(organization_id=org.id, owner_user_id=actor.id, title="Backend Engineer", requirements="Python FastAPI 3 years of experience")
    db.add_all([candidate, job])
    db.flush()
    application = create_application(db, candidate=candidate, job=job, owner_user_id=actor.id)
    db.commit()
    response = client.post(f"/api/jobs/{job.id}/match?threshold=80")
    assert response.status_code == 200
    results = response.json()["results"]
    assert len(results) == 1
    assert results[0]["match_score"] == 100
    assert "Skills fit" in results[0]["explanation"]
    db.refresh(application)
    assert application.match_metadata["provider"] == "local_rule_engine"
    assert application.match_metadata["method"] == "rule"
